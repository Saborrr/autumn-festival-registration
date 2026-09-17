from io import BytesIO

from django.http import HttpResponse
from django.utils import timezone
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from .models import Registration


HEADERS = (
    "№",
    "Участник",
    "Формат участия",
    "Команда",
    "Подразделение",
    "Телефон",
    "Возраст",
    "Сопровождение",
    "Статус",
    "Дата регистрации",
)


def registration_rows():
    registrations = Registration.objects.select_related("team").prefetch_related(
        "companion_children"
    ).order_by(
        "team__number", "created_at", "pk"
    )
    rows = []
    for number, registration in enumerate(registrations, start=1):
        age = registration.get_participant_type_display() or "—"
        if registration.child_age is not None:
            age = f"{age}, {registration.child_age} лет"
        companions = registration.get_companions_display()
        child_summary = registration.companion_children_display
        if child_summary:
            companions = f"{companions}: {child_summary}"
        if registration.children_details:
            companions = f"{companions}; архив: {registration.children_details}"
        rows.append(
            (
                number,
                registration.full_name,
                registration.participation_display,
                f"Команда {registration.team.number}" if registration.team else "—",
                registration.organization_unit or "—",
                registration.phone or "—",
                age,
                companions,
                registration.get_status_display(),
                timezone.localtime(registration.created_at).strftime("%d.%m.%Y %H:%M"),
            )
        )
    return rows


def xlsx_response():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Регистрации"
    sheet.append(HEADERS)
    for row in registration_rows():
        sheet.append(row)

    header_fill = PatternFill("solid", fgColor="7A2E16")
    header_font = Font(color="FFFFFF", bold=True)
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    widths = (6, 28, 42, 14, 24, 20, 18, 42, 14, 19)
    for column, width in zip(sheet.columns, widths):
        sheet.column_dimensions[column[0].column_letter].width = width
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.row_dimensions[1].height = 32
    sheet.sheet_view.showGridLines = False
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.sheet_properties.pageSetUpPr.fitToPage = True

    output = BytesIO()
    workbook.save(output)
    response = HttpResponse(
        output.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = 'attachment; filename="registrations.xlsx"'
    response["Cache-Control"] = "private, no-store"
    return response


def docx_response():
    document = Document()
    section = document.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = section.page_height, section.page_width
    section.top_margin = Cm(1.2)
    section.bottom_margin = Cm(1.2)
    section.left_margin = Cm(1.2)
    section.right_margin = Cm(1.2)

    title = document.add_heading("Краски осени — результаты регистрации", level=1)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    generated = document.add_paragraph(
        f"Сформировано: {timezone.localtime().strftime('%d.%m.%Y %H:%M')}"
    )
    generated.alignment = WD_ALIGN_PARAGRAPH.CENTER

    rows = registration_rows()
    table = document.add_table(rows=1, cols=len(HEADERS))
    table.style = "Table Grid"
    table.autofit = True
    for index, header in enumerate(HEADERS):
        cell = table.rows[0].cells[index]
        cell.text = header
        for run in cell.paragraphs[0].runs:
            run.bold = True
            run.font.size = Pt(7)

    for row in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            cells[index].text = str(value)
            for paragraph in cells[index].paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(7)

    output = BytesIO()
    document.save(output)
    response = HttpResponse(
        output.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )
    response["Content-Disposition"] = 'attachment; filename="registrations.docx"'
    response["Cache-Control"] = "private, no-store"
    return response
