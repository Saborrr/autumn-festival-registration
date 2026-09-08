from io import BytesIO

from django.contrib.auth import get_user_model
from django.test import TestCase
from django.urls import reverse
from docx import Document
from openpyxl import load_workbook

from registrations.models import Registration, Team


class HomePageTests(TestCase):
    def test_home_lists_all_ten_teams_with_capacity(self):
        response = self.client.get(reverse("home"))

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Краски осени")
        self.assertEqual(len(response.context["teams"]), 10)
        self.assertContains(response, "0 из 10")
        self.assertContains(response, "19 сентября")
        self.assertContains(
            response,
            "За подробной информацией обращайтесь к ЛРА",
        )
        self.assertNotContains(response, "За более подробной информацией")
        self.assertContains(response, "Участник «Большой игры»")
        self.assertContains(response, "Болельщик")
        self.assertContains(response, "Конкурс «Счастье в пирогах»")
        self.assertContains(
            response,
            "В каждую команду желательно включить двух детей младше 14 лет",
        )
        self.assertNotContains(response, "желательно предусмотреть двух детей")
        self.assertContains(
            response,
            "Дизайн и разработка Александр Фадеев, al.al.fdv@gmail.com",
        )
        self.assertNotContains(response, "Фото:")

    def test_home_shows_active_participants_in_their_team(self):
        team = Team.objects.get(number=1)
        team.registrations.create(
            full_name="Мария Иванова", participant_type="adult"
        )
        team.registrations.create(
            full_name="Павел Орлов",
            participant_type="adult",
            status="cancelled",
        )

        response = self.client.get(reverse("home"))

        self.assertContains(response, "Уже в команде")
        self.assertContains(response, "Мария Иванова")
        self.assertNotContains(response, "Павел Орлов")

    def test_valid_registration_reserves_a_place(self):
        team = Team.objects.get(number=1)

        response = self.client.post(
            reverse("register"),
            {
                "team": team.pk,
                "full_name": "Мария Иванова",
                "organization_unit": "ЭФКО, производство",
                "phone": "+7 900 000-00-00",
                "activities": ["game"],
                "participant_type": "adult",
                "child_age": "",
                "companions": "alone",
                "children_details": "",
            },
        )

        self.assertRedirects(response, reverse("registration_success"))
        self.assertTrue(
            team.registrations.filter(full_name="Мария Иванова").exists()
        )

    def test_registration_is_rejected_when_team_is_full(self):
        team = Team.objects.get(number=1)
        for index in range(10):
            team.registrations.create(
                full_name=f"Участник {index}", participant_type="adult"
            )

        response = self.client.post(
            reverse("register"),
            {
                "team": team.pk,
                "full_name": "Одиннадцатый Участник",
                "organization_unit": "ЭФКО, производство",
                "phone": "+7 900 000-00-00",
                "activities": ["game"],
                "participant_type": "adult",
                "child_age": "",
                "companions": "alone",
                "children_details": "",
            },
        )

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "В этой команде уже нет свободных мест")
        self.assertEqual(team.registrations.count(), 10)

    def test_registration_page_uses_category_neutral_copy(self):
        response = self.client.get(reverse("register") + "?format=pie")

        self.assertContains(response, "Ваш формат")
        self.assertContains(response, "Вернуться к форматам")
        self.assertNotContains(response, "Ваше место<br>в команде", html=True)

    def test_companion_options_are_checkboxes(self):
        response = self.client.get(reverse("register"))

        self.assertContains(response, 'name="companions"', count=3)
        self.assertContains(response, 'type="checkbox"', count=6)
        self.assertContains(response, "Со мной будет 1 взрослый родственник")
        self.assertContains(response, "Со мной будет ребёнок (дети)")

    def test_adult_relative_and_children_can_be_selected_together(self):
        response = self.client.post(
            reverse("register"),
            {
                "team": "",
                "full_name": "Семейный Болельщик",
                "organization_unit": "ЭФКО, офис",
                "phone": "+7 900 555-66-77",
                "activities": ["supporter"],
                "participant_type": "",
                "child_age": "",
                "companions": ["adult_relative", "children"],
                "children_details": "2 ребёнка, 6 и 9 лет",
            },
        )

        self.assertRedirects(response, reverse("registration_success"))
        registration = Registration.objects.get(full_name="Семейный Болельщик")
        self.assertEqual(registration.companions, "adult_and_children")
        self.assertEqual(registration.children_details, "2 ребёнка, 6 и 9 лет")

    def test_alone_cannot_be_combined_with_companions(self):
        response = self.client.post(
            reverse("register"),
            {
                "team": "",
                "full_name": "Противоречивый Выбор",
                "organization_unit": "ЭФКО, офис",
                "phone": "+7 900 555-66-88",
                "activities": ["supporter"],
                "participant_type": "",
                "child_age": "",
                "companions": ["alone", "adult_relative"],
                "children_details": "",
            },
        )

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Вариант «Приду один (одна)» нельзя совмещать")
        self.assertFalse(
            Registration.objects.filter(full_name="Противоречивый Выбор").exists()
        )

    def test_combined_companions_require_children_details(self):
        response = self.client.post(
            reverse("register"),
            {
                "team": "",
                "full_name": "Без Возраста Детей",
                "organization_unit": "ЭФКО, офис",
                "phone": "+7 900 555-66-99",
                "activities": ["supporter"],
                "participant_type": "",
                "child_age": "",
                "companions": ["adult_relative", "children"],
                "children_details": "",
            },
        )

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Укажите количество и возраст детей")
        self.assertFalse(Registration.objects.filter(full_name="Без Возраста Детей").exists())

    def test_supporter_can_register_without_a_team(self):
        response = self.client.post(
            reverse("register"),
            {
                "team": "",
                "full_name": "Елена Соколова",
                "organization_unit": "ЭФКО, офис",
                "phone": "+7 953 111-22-33",
                "activities": ["supporter"],
                "participant_type": "",
                "child_age": "",
                "companions": "adult_relative",
                "children_details": "",
            },
        )

        self.assertRedirects(response, reverse("registration_success"))
        registration = Team.objects.first().registrations.model.objects.get(
            full_name="Елена Соколова"
        )
        self.assertEqual(registration.participation_format, "supporter")
        self.assertIsNone(registration.team)

    def test_pie_contest_is_a_separate_registration_category(self):
        response = self.client.post(
            reverse("register"),
            {
                "team": "",
                "full_name": "Ольга Пирогова",
                "organization_unit": "ЭФКО, логистика",
                "phone": "+7 953 444-55-66",
                "activities": ["pie"],
                "participant_type": "",
                "child_age": "",
                "companions": "children",
                "children_details": "2 ребёнка, 6 и 9 лет",
            },
        )

        self.assertRedirects(response, reverse("registration_success"))
        registration = Registration.objects.get(full_name="Ольга Пирогова")
        self.assertTrue(registration.participates_in_pie)

    def test_game_participant_can_also_enter_pie_contest(self):
        team = Team.objects.get(number=2)
        response = self.client.post(
            reverse("register"),
            {
                "team": team.pk,
                "full_name": "Иван Игров",
                "organization_unit": "ТПК",
                "phone": "+7 900 111-22-33",
                "activities": ["game", "pie"],
                "participant_type": "adult",
                "child_age": "",
                "companions": "alone",
                "children_details": "",
            },
        )

        self.assertRedirects(response, reverse("registration_success"))
        registration = Registration.objects.get(full_name="Иван Игров")
        self.assertEqual(registration.participation_format, "game")
        self.assertTrue(registration.participates_in_pie)
        self.assertEqual(registration.team, team)

    def test_supporter_can_also_enter_pie_contest(self):
        response = self.client.post(
            reverse("register"),
            {
                "team": "",
                "full_name": "Анна Болельщикова",
                "organization_unit": "УК",
                "phone": "+7 900 222-33-44",
                "activities": ["supporter", "pie"],
                "participant_type": "",
                "child_age": "",
                "companions": "alone",
                "children_details": "",
            },
        )

        self.assertRedirects(response, reverse("registration_success"))
        registration = Registration.objects.get(full_name="Анна Болельщикова")
        self.assertEqual(registration.participation_format, "supporter")
        self.assertTrue(registration.participates_in_pie)
        self.assertIsNone(registration.team)

    def test_game_participant_cannot_also_be_supporter(self):
        response = self.client.post(
            reverse("register"),
            {
                "team": Team.objects.get(number=2).pk,
                "full_name": "Неверный Выбор",
                "organization_unit": "УК",
                "phone": "+7 900 333-44-55",
                "activities": ["game", "supporter"],
                "participant_type": "adult",
                "child_age": "",
                "companions": "alone",
                "children_details": "",
            },
        )

        self.assertEqual(response.status_code, 200)
        self.assertContains(
            response,
            "Нельзя одновременно выбрать участие в «Большой игре» и роль болельщика",
        )
        self.assertFalse(Registration.objects.filter(full_name="Неверный Выбор").exists())


class ModeratorDashboardTests(TestCase):
    def test_companion_filters_include_combined_registration(self):
        moderator = get_user_model().objects.create_user(
            username="family-moderator", password="strong-password", is_staff=True
        )
        self.client.force_login(moderator)
        combined = Registration.objects.create(
            full_name="Семейный участник",
            participation_format="supporter",
            companions="adult_and_children",
            children_details="1 ребёнок, 8 лет",
        )

        for companion_filter in ("adult_relative", "children"):
            with self.subTest(companion_filter=companion_filter):
                response = self.client.get(
                    reverse("moderator_dashboard"),
                    {"companions": companion_filter},
                )
                self.assertEqual(list(response.context["registrations"]), [combined])

    def test_dashboard_displays_and_filters_combined_pie_participation(self):
        moderator = get_user_model().objects.create_user(
            username="combined-moderator", password="strong-password", is_staff=True
        )
        self.client.force_login(moderator)
        combined = Team.objects.get(number=1).registrations.create(
            full_name="Игрок и пекарь",
            participation_format="game",
            participates_in_pie=True,
            participant_type="adult",
        )

        response = self.client.get(reverse("moderator_dashboard"), {"format": "pie"})

        self.assertEqual(list(response.context["registrations"]), [combined])
        self.assertContains(
            response,
            "Участник «Большой игры» + Конкурс «Счастье в пирогах»",
        )

    def test_dashboard_requires_staff_and_shows_registrations(self):
        team = Team.objects.get(number=3)
        team.registrations.create(
            full_name="Елена Соколова", participant_type="adult"
        )

        anonymous = self.client.get(reverse("moderator_dashboard"))
        self.assertRedirects(
            anonymous,
            f'{reverse("login")}?next={reverse("moderator_dashboard")}',
        )

        moderator = get_user_model().objects.create_user(
            username="moderator", password="strong-password", is_staff=True
        )
        self.client.force_login(moderator)
        response = self.client.get(reverse("moderator_dashboard"))

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Елена Соколова")
        self.assertContains(response, "Команда 3")

    def test_staff_can_cancel_registration(self):
        team = Team.objects.get(number=4)
        registration = team.registrations.create(
            full_name="Павел Орлов", participant_type="adult"
        )
        moderator = get_user_model().objects.create_user(
            username="moderator", password="strong-password", is_staff=True
        )
        self.client.force_login(moderator)

        response = self.client.post(
            reverse("moderator_toggle", args=[registration.pk])
        )

        self.assertRedirects(response, reverse("moderator_dashboard"))
        registration.refresh_from_db()
        self.assertEqual(registration.status, registration.Status.CANCELLED)

    def test_staff_cannot_restore_registration_into_full_team(self):
        team = Team.objects.get(number=5)
        cancelled = team.registrations.create(
            full_name="Запасной Участник",
            participant_type="adult",
            status="cancelled",
        )
        for index in range(10):
            team.registrations.create(
                full_name=f"Активный {index}", participant_type="adult"
            )
        moderator = get_user_model().objects.create_user(
            username="moderator", password="strong-password", is_staff=True
        )
        self.client.force_login(moderator)

        response = self.client.post(reverse("moderator_toggle", args=[cancelled.pk]))

        self.assertRedirects(response, reverse("moderator_dashboard"))
        cancelled.refresh_from_db()
        self.assertEqual(cancelled.status, cancelled.Status.CANCELLED)

    def test_dashboard_filters_every_supported_column(self):
        moderator = get_user_model().objects.create_user(
            username="filter-moderator", password="strong-password", is_staff=True
        )
        self.client.force_login(moderator)
        team = Team.objects.get(number=2)
        game = team.registrations.create(
            full_name="Игрок",
            organization_unit="Производство",
            phone="111",
            participation_format="game",
            participant_type="adult",
            companions="alone",
        )
        supporter = Registration.objects.create(
            full_name="Болельщик",
            organization_unit="Офис",
            phone="222",
            participation_format="supporter",
            participant_type="",
            companions="adult_relative",
        )
        pie = Registration.objects.create(
            full_name="Пекарь",
            organization_unit="Логистика",
            phone="333",
            participation_format="pie",
            participant_type="",
            companions="children",
            children_details="1 ребёнок, 8 лет",
            status="cancelled",
        )
        cases = (
            ({"format": "supporter"}, supporter),
            ({"team": str(team.pk)}, game),
            ({"organization_unit": "Логистика"}, pie),
            ({"companions": "alone"}, game),
            ({"status": "cancelled"}, pie),
        )

        for params, expected in cases:
            with self.subTest(params=params):
                response = self.client.get(reverse("moderator_dashboard"), params)
                self.assertEqual(list(response.context["registrations"]), [expected])

    def test_dashboard_sorts_every_supported_column(self):
        moderator = get_user_model().objects.create_user(
            username="sort-moderator", password="strong-password", is_staff=True
        )
        self.client.force_login(moderator)
        team_one = Team.objects.get(number=1)
        team_three = Team.objects.get(number=3)
        alpha = team_one.registrations.create(
            full_name="Альфа",
            organization_unit="Альфа",
            participation_format="game",
            participant_type="adult",
            companions="adult_relative",
            status="active",
        )
        omega = team_three.registrations.create(
            full_name="Омега",
            organization_unit="Омега",
            participation_format="pie",
            participant_type="",
            companions="children",
            children_details="1 ребёнок, 8 лет",
            status="cancelled",
        )
        cases = (
            ("format", "asc", alpha),
            ("team", "desc", omega),
            ("organization_unit", "asc", alpha),
            ("companions", "asc", alpha),
            ("status", "desc", omega),
        )

        for sort, direction, expected_first in cases:
            with self.subTest(sort=sort, direction=direction):
                response = self.client.get(
                    reverse("moderator_dashboard"),
                    {"sort": sort, "direction": direction},
                )
                self.assertEqual(response.context["registrations"][0], expected_first)

    def test_dashboard_does_not_offer_participant_or_phone_sorting(self):
        moderator = get_user_model().objects.create_user(
            username="safe-moderator", password="strong-password", is_staff=True
        )
        self.client.force_login(moderator)

        response = self.client.get(reverse("moderator_dashboard"))

        self.assertNotContains(response, 'value="full_name"')
        self.assertNotContains(response, 'value="phone"')


class RegistrationExportTests(TestCase):
    def setUp(self):
        self.moderator = get_user_model().objects.create_user(
            username="export-moderator", password="strong-password", is_staff=True
        )
        self.registration = Team.objects.get(number=2).registrations.create(
            full_name="Анна Экспортова",
            organization_unit="Производство",
            phone="+7 900 123-45-67",
            participation_format="game",
            participates_in_pie=True,
            participant_type="adult",
            companions="alone",
        )

    def test_exports_require_staff_access(self):
        for name in ("export_registrations_xlsx", "export_registrations_docx"):
            with self.subTest(name=name):
                response = self.client.get(reverse(name))
                self.assertRedirects(response, f'{reverse("login")}?next={reverse(name)}')

    def test_staff_can_download_excel_registration_list(self):
        self.client.force_login(self.moderator)

        response = self.client.get(reverse("export_registrations_xlsx"))

        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response["Content-Type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        self.assertIn("registrations.xlsx", response["Content-Disposition"])
        self.assertEqual(response["Cache-Control"], "private, no-store")
        workbook = load_workbook(BytesIO(response.content), read_only=True)
        rows = list(workbook["Регистрации"].values)
        self.assertEqual(rows[0][1], "Участник")
        self.assertEqual(rows[1][1], "Анна Экспортова")
        self.assertIn("Счастье в пирогах", rows[1][2])

    def test_staff_can_download_word_registration_list(self):
        self.client.force_login(self.moderator)

        response = self.client.get(reverse("export_registrations_docx"))

        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response["Content-Type"],
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
        self.assertIn("registrations.docx", response["Content-Disposition"])
        self.assertEqual(response["Cache-Control"], "private, no-store")
        document = Document(BytesIO(response.content))
        cells = [cell.text for table in document.tables for row in table.rows for cell in row.cells]
        self.assertIn("Анна Экспортова", cells)
        self.assertTrue(any("Счастье в пирогах" in cell for cell in cells))

    def test_exports_show_combined_companions(self):
        self.registration.companions = "adult_and_children"
        self.registration.children_details = "1 ребёнок, 7 лет"
        self.registration.save()
        self.client.force_login(self.moderator)

        response = self.client.get(reverse("export_registrations_xlsx"))

        workbook = load_workbook(BytesIO(response.content), read_only=True)
        rows = list(workbook["Регистрации"].values)
        self.assertIn("1 взрослый родственник", rows[1][7])
        self.assertIn("ребёнок (дети)", rows[1][7])
        self.assertIn("1 ребёнок, 7 лет", rows[1][7])

    def test_export_buttons_are_available_in_moderator_and_admin_panels(self):
        self.client.force_login(self.moderator)
        moderator_response = self.client.get(reverse("moderator_dashboard"))
        self.assertContains(moderator_response, reverse("export_registrations_xlsx"))
        self.assertContains(moderator_response, reverse("export_registrations_docx"))

        administrator = get_user_model().objects.create_superuser(
            username="export-admin", password="strong-password"
        )
        self.client.force_login(administrator)
        admin_response = self.client.get(
            reverse("admin:registrations_registration_changelist")
        )
        self.assertContains(admin_response, reverse("export_registrations_xlsx"))
        self.assertContains(admin_response, reverse("export_registrations_docx"))
